#!/usr/bin/env python3
"""Generate the committed inspection template catalog from the appraisal workbook.

This script is a development/import tool only. The application and seed consume the
generated JSON and never read the external workbook at runtime.
"""

from __future__ import annotations

import argparse
import json
import re
from collections import OrderedDict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


POSITION_SHEETS = OrderedDict(
    [
        ("左前方位", ("FRONT_LEFT", "左前", "LEFT", "LF")),
        ("左后方位", ("REAR_LEFT", "左后", "LEFT", "LR")),
        ("右前方位", ("FRONT_RIGHT", "右前", "RIGHT", "RF")),
        ("右后方位", ("REAR_RIGHT", "右后", "RIGHT", "RR")),
    ]
)


def clean(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\n", " ")).strip()


def is_empty_or_dash(value: Any) -> bool:
    text = clean(value)
    return not text or all(char in "—–-" for char in text)


def split_phenomena(value: Any) -> list[str]:
    text = clean(value)
    if is_empty_or_dash(text):
        return []
    pieces = re.split(r"[、；;/]+", text)
    result: list[str] = []
    for piece in pieces:
        item = piece.strip(" ，,。；;、/")
        if item and not is_empty_or_dash(item) and item not in result:
            result.append(item)
    return result or [text]


def code_fragment(text: str) -> str:
    normalized = re.sub(r"[^A-Za-z0-9]+", "_", text.upper()).strip("_")
    return normalized or "ITEM"


def normalized_component(value: Any) -> str:
    return clean(value).replace("★", "").strip()


def component_class(component: str, group: str, has_accident_column: bool, starred: bool, exterior_only: bool) -> str:
    if starred or any(token in component for token in ("纵梁", "横梁", "柱", "底板", "框架", "边梁", "门槛", "围板", "轮旋", "减震器座", "骨架", "大顶")):
        return "STRUCTURAL"
    if "防撞" in component or "加强" in component:
        return "REINFORCEMENT"
    if exterior_only:
        return "COVERAGE"
    if has_accident_column and "结构" in group:
        return "STRUCTURAL"
    return "OTHER"


def criterion(
    *,
    code: str,
    label: str,
    axis: str,
    criterion_type: str,
    source_sheet: str,
    source_row: int,
    source_column: str,
    source_text: str,
    source_note: str | None = None,
    source_part_index: int = 1,
    rule_key: str | None = None,
    counts_flood: bool = False,
    hard_stop: bool = False,
) -> dict[str, Any]:
    return {
        "code": code,
        "label": label,
        "axis": axis,
        "criterionType": criterion_type,
        "description": source_text,
        "ruleKey": rule_key,
        "countsAsDistinctFloodFinding": counts_flood,
        "hardStopCriterion": hard_stop,
        "sourceSheet": source_sheet,
        "sourceRow": source_row,
        "sourceColumn": source_column,
        "sourcePartIndex": source_part_index,
        "sourceText": source_text,
        "sourceNote": source_note,
    }


def section(code: str, name: str, axis: str, description: str, positions: list[dict[str, Any]]) -> dict[str, Any]:
    return {"code": code, "name": name, "axis": axis, "description": description, "positions": positions}


def position(code: str, name: str, *, side: str | None = None, flood_key: str | None = None, items: list[dict[str, Any]]) -> dict[str, Any]:
    return {"code": code, "name": name, "side": side, "floodAggregationKey": flood_key, "checkItems": items}


def check_item(
    *,
    code: str,
    name: str,
    category: str,
    component: str,
    group: str,
    source_sheet: str,
    source_row: int,
    source_section: str,
    criteria: list[dict[str, Any]],
    accident_participant: bool = False,
    component_class: str = "OTHER",
) -> dict[str, Any]:
    return {
        "code": code,
        "name": name,
        "category": category,
        "componentClass": component_class,
        "accidentDecisionParticipant": accident_participant,
        "description": f"来源于{source_sheet}第{source_row}行：{name}",
        "sourceSheet": source_sheet,
        "sourceRow": source_row,
        "sourceSection": source_section,
        "sourceText": component,
        "criteria": criteria,
    }


def sheet_header(ws: Any) -> list[str]:
    return [clean(ws.cell(2, column).value) for column in range(1, ws.max_column + 1)]


def build_position_sheet(ws: Any, sheet_name: str) -> tuple[dict[str, Any], dict[str, Any]]:
    position_code, position_name, side, prefix = POSITION_SHEETS[sheet_name]
    current_section = "未分组"
    items: list[dict[str, Any]] = []
    expected_criteria = 0
    fallback_criteria = 0
    valid_rows = 0
    source_criteria: list[dict[str, Any]] = []
    headers = {"B": "ACCIDENT_STRUCTURE", "C": "FLOOD_DAMAGE", "D": "EXTERIOR", "E": "INTERIOR"}
    axis_by_column = {"B": ("ACCIDENT_HISTORY", "ACCIDENT_DEFECT"), "C": ("FLOOD_DAMAGE", "FLOOD_FINDING"), "D": ("EXTERIOR_INTERIOR", "EXTERIOR_FINDING"), "E": ("EXTERIOR_INTERIOR", "INTERIOR_FINDING")}

    for row in range(3, ws.max_row + 1):
        component_raw = clean(ws.cell(row, 1).value)
        if component_raw.startswith("【"):
            current_section = component_raw.strip("【】")
            continue
        if not component_raw:
            continue
        valid_rows += 1
        component = normalized_component(component_raw)
        row_criteria: list[dict[str, Any]] = []
        category = "OTHER"
        has_accident = False
        non_empty_axes: list[str] = []
        for column, (axis, criterion_type) in axis_by_column.items():
            raw_text = clean(ws[f"{column}{row}"].value)
            if is_empty_or_dash(raw_text):
                continue
            non_empty_axes.append(axis)
            if axis == "ACCIDENT_HISTORY":
                category = "STRUCTURE"
                has_accident = True
            elif axis == "FLOOD_DAMAGE" and category == "OTHER":
                category = "ELECTRICAL"
            elif axis == "EXTERIOR_INTERIOR" and category == "OTHER":
                category = "EXTERIOR" if criterion_type == "EXTERIOR_FINDING" else "CABIN"
            pieces = split_phenomena(raw_text)
            for index, piece in enumerate(pieces, start=1):
                criterion_code = f"{criterion_type[:4]}_{index:02d}"
                row_criteria.append(
                    criterion(
                        code=criterion_code,
                        label=piece,
                        axis=axis,
                        criterion_type=criterion_type,
                        source_sheet=sheet_name,
                        source_row=row,
                        source_column=column,
                        source_text=raw_text,
                        source_part_index=index,
                        rule_key="ACCIDENT_COMPONENT_DEFECT" if axis == "ACCIDENT_HISTORY" else "FLOOD_DISTINCT_FINDING" if axis == "FLOOD_DAMAGE" else "POSITIONAL_OBSERVATION",
                        counts_flood=axis == "FLOOD_DAMAGE",
                    )
                )
                source_criteria.append({"row": row, "column": column, "partIndex": index, "text": raw_text})
                expected_criteria += 1
        if not row_criteria:
            row_criteria.append(
                criterion(
                    code="OBS_01",
                    label="现场观察异常",
                    axis="EXTERIOR_INTERIOR",
                    criterion_type="OBSERVATION",
                    source_sheet=sheet_name,
                    source_row=row,
                    source_column="F",
                    source_text="原表未列具体现象；保留现场观察结果",
                    source_part_index=1,
                    rule_key="POSITIONAL_OBSERVATION",
                )
            )
            source_criteria.append({"row": row, "column": "F", "partIndex": 1, "text": "原表未列具体现象；保留现场观察结果"})
            expected_criteria += 1
            fallback_criteria += 1
        starred = "★" in component_raw
        exterior_only = bool(row_criteria) and all(item["axis"] == "EXTERIOR_INTERIOR" for item in row_criteria)
        item = check_item(
            code=f"{prefix}_R{row:02d}",
            name=component,
            category=category,
            component=component_raw,
            group=current_section,
            source_sheet=sheet_name,
            source_row=row,
            source_section=current_section,
            criteria=row_criteria,
            accident_participant=starred,
            component_class=component_class(component, current_section, has_accident, starred, exterior_only),
        )
        items.append(item)

    section_data = section(
        f"POSITION_{position_code}",
        f"{position_name}方位（来源模板）",
        "POSITIONAL_INSPECTION",
        "来自方位式鉴定表；一个组件可同时拥有事故、水泡、外观和内饰标准。",
        [position(position_code, position_name, side=side, flood_key=side, items=items)],
    )
    summary = {
        "sheet": sheet_name,
        "headers": sheet_header(ws),
        "validRows": valid_rows,
        "checkItems": len(items),
        "criteria": expected_criteria,
        "fallbackCriteria": fallback_criteria,
        "sourceCriteria": source_criteria,
    }
    return section_data, summary


def build_general_sheet(ws: Any) -> tuple[dict[str, Any], dict[str, Any]]:
    current_section = "未分组"
    items: list[dict[str, Any]] = []
    source_criteria: list[dict[str, Any]] = []
    for row in range(3, ws.max_row + 1):
        group = clean(ws.cell(row, 1).value)
        if group.startswith("【"):
            current_section = group.strip("【】")
            continue
        name = clean(ws.cell(row, 2).value)
        if not name:
            continue
        content = clean(ws.cell(row, 3).value)
        note = clean(ws.cell(row, 4).value)
        label = content or "现场核对"
        item_criteria = [
            criterion(
                code="GENERAL_01",
                label=label,
                axis="LEGAL_TRADEABILITY" if "手续" in current_section else "DISCLOSURE",
                criterion_type="HARD_STOP" if "停止" in note else "GENERAL_CHECK",
                source_sheet="整车通用",
                source_row=row,
                source_column="C",
                source_text=content,
                source_note=note or None,
                hard_stop="停止" in note,
                rule_key="LEGAL_HARD_STOP" if "停止" in note else "GENERAL_DISCLOSURE",
            )
        ]
        source_criteria.append({"row": row, "column": "C", "partIndex": 1, "text": content})
        items.append(
            check_item(
                code=f"GENERAL_R{row:02d}",
                name=name,
                category="DOCUMENT",
                component=name,
                group=current_section,
                source_sheet="整车通用",
                source_row=row,
                source_section=current_section,
                criteria=item_criteria,
                component_class="OTHER",
            )
        )
    section_data = section(
        "LEGAL_TRADEABILITY",
        "手续与可交易性",
        "LEGAL_TRADEABILITY",
        "来自整车通用表；关键资料不一致时停止鉴定或交易流转。",
        [position("GENERAL", "整车通用", items=items)],
    )
    return section_data, {"sheet": "整车通用", "headers": sheet_header(ws), "validRows": len(items), "checkItems": len(items), "criteria": len(source_criteria), "fallbackCriteria": 0, "sourceCriteria": source_criteria}


def build_powertrain_sheet(ws: Any) -> tuple[dict[str, Any], dict[str, Any]]:
    group_to_position = {"发动机（机舱静态）": ("ENGINE", "发动机", "CORE_FUNCTION", "POWERTRAIN"), "变速器": ("TRANSMISSION", "变速器", "CORE_FUNCTION", "POWERTRAIN"), "底盘": ("CHASSIS", "底盘", "CURRENT_SAFETY", "CHASSIS")}
    current_group = ""
    positions: OrderedDict[str, dict[str, Any]] = OrderedDict()
    source_criteria: list[dict[str, Any]] = []
    for row in range(3, ws.max_row + 1):
        group = clean(ws.cell(row, 1).value)
        if group.startswith("【"):
            current_group = group.strip("【】")
            continue
        name = clean(ws.cell(row, 2).value)
        if not name:
            continue
        position_code, position_name, axis, category = group_to_position.get(current_group, ("MECHANICAL", "机械专项", "CORE_FUNCTION", "OTHER"))
        raw_text = clean(ws.cell(row, 3).value)
        pieces = split_phenomena(raw_text)
        item_criteria = []
        for index, piece in enumerate(pieces, start=1):
            criterion_type = "SAFETY_FINDING" if axis == "CURRENT_SAFETY" else "FUNCTION_FINDING"
            item_criteria.append(
                criterion(
                    code=f"{criterion_type[:4]}_{index:02d}",
                    label=piece,
                    axis=axis,
                    criterion_type=criterion_type,
                    source_sheet="三大件专项",
                    source_row=row,
                    source_column="C",
                    source_text=raw_text,
                    source_note=clean(ws.cell(row, 4).value) or None,
                    source_part_index=index,
                    rule_key="CURRENT_SAFETY_GATE" if axis == "CURRENT_SAFETY" else "CORE_FUNCTION_GATE",
                )
            )
            source_criteria.append({"row": row, "column": "C", "partIndex": index, "text": raw_text})
        positions.setdefault(position_code, position(position_code, position_name, items=[]))
        positions[position_code]["checkItems"].append(
            check_item(
                code=f"POWER_R{row:02d}",
                name=name,
                category=category,
                component=name,
                group=current_group,
                source_sheet="三大件专项",
                source_row=row,
                source_section=current_group,
                criteria=item_criteria,
                component_class="OTHER",
            )
        )
    section_data = section("MECHANICAL_SYSTEMS", "三大件专项", "CORE_FUNCTION", "发动机、变速器和底盘专项；底盘条目通过 criterion.axis 进入当前安全结论。", list(positions.values()))
    return section_data, {"sheet": "三大件专项", "headers": sheet_header(ws), "validRows": sum(len(item["checkItems"]) for item in positions.values()), "checkItems": sum(len(item["checkItems"]) for item in positions.values()), "criteria": len(source_criteria), "fallbackCriteria": 0, "sourceCriteria": source_criteria}


def build_new_energy_sheet(ws: Any) -> tuple[dict[str, Any], dict[str, Any]]:
    group_to_position = {
        "三电系统外观": ("NEW_ENERGY_APPEARANCE", "三电系统外观"),
        "电池性能（需设备读取）": ("NEW_ENERGY_BATTERY_HEALTH", "电池性能"),
        "电机运转": ("NEW_ENERGY_MOTOR_OPERATION", "电机运转"),
    }
    current_group = ""
    positions: OrderedDict[str, dict[str, Any]] = OrderedDict()
    source_criteria: list[dict[str, Any]] = []
    for row in range(3, ws.max_row + 1):
        group = clean(ws.cell(row, 1).value)
        if group.startswith("【"):
            current_group = group.strip("【】")
            continue
        name = clean(ws.cell(row, 2).value)
        if not name:
            continue
        position_code, position_name = group_to_position.get(current_group, ("NEW_ENERGY_SYSTEM", "新能源系统"))
        raw_text = clean(ws.cell(row, 3).value)
        pieces = split_phenomena(raw_text)
        item_criteria = []
        for index, piece in enumerate(pieces, start=1):
            criterion_type = "NEW_ENERGY_MEASURED" if "需设备" in current_group or any(token in name.upper() for token in ("SOH", "压差", "循环", "故障码", "绝缘")) else "NEW_ENERGY_FINDING"
            item_criteria.append(
                criterion(
                    code=f"{criterion_type[:4]}_{index:02d}",
                    label=piece,
                    axis="NEW_ENERGY",
                    criterion_type=criterion_type,
                    source_sheet="新能源专项",
                    source_row=row,
                    source_column="C",
                    source_text=raw_text,
                    source_note=clean(ws.cell(row, 4).value) or None,
                    source_part_index=index,
                    rule_key="NEW_ENERGY_HEALTH" if criterion_type == "NEW_ENERGY_MEASURED" else "NEW_ENERGY_SAFETY",
                )
            )
            source_criteria.append({"row": row, "column": "C", "partIndex": index, "text": raw_text})
        positions.setdefault(position_code, position(position_code, position_name, items=[]))
        positions[position_code]["checkItems"].append(
            check_item(
                code=f"NE_R{row:02d}",
                name=name,
                category="NEW_ENERGY",
                component=name,
                group=current_group,
                source_sheet="新能源专项",
                source_row=row,
                source_section=current_group,
                criteria=item_criteria,
                component_class="OTHER",
            )
        )
    section_data = section("NEW_ENERGY", "新能源专项", "NEW_ENERGY", "所有新能源车辆共用此版本；保留设备读数和高压安全结论。", list(positions.values()))
    return section_data, {"sheet": "新能源专项", "headers": sheet_header(ws), "validRows": sum(len(item["checkItems"]) for item in positions.values()), "checkItems": sum(len(item["checkItems"]) for item in positions.values()), "criteria": len(source_criteria), "fallbackCriteria": 0, "sourceCriteria": source_criteria}


def build_rules(ws: Any) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    rules: list[dict[str, Any]] = []
    excluded: list[dict[str, Any]] = []
    for row in range(2, ws.max_row + 1):
        label = clean(ws.cell(row, 1).value)
        text = clean(ws.cell(row, 2).value)
        if not label or not text or label == "签名栏":
            continue
        if label.startswith("③") or "火烧" in label:
            excluded.append({"sheet": "判定总则", "row": row, "label": label, "reason": "fire-damage rules are out of current scope", "sourceText": text})
            continue
        code = {"① 事故车": "ACCIDENT_RULE", "② 泡水车": "FLOOD_RULE", "④ 其他": "BOUNDARY_RULE"}.get(label, f"GENERAL_RULE_{row:02d}")
        rules.append({"code": code, "label": label, "sourceSheet": "判定总则", "sourceRow": row, "sourceText": text, "scope": "ACTIVE"})
    return rules, excluded


def build_catalog(workbook_path: Path) -> dict[str, Any]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    standard_sections: list[dict[str, Any]] = []
    summaries: list[dict[str, Any]] = []
    general, summary = build_general_sheet(workbook["整车通用"])
    standard_sections.append(general)
    summaries.append(summary)
    for sheet_name in POSITION_SHEETS:
        physical, summary = build_position_sheet(workbook[sheet_name], sheet_name)
        standard_sections.append(physical)
        summaries.append(summary)
    power, summary = build_powertrain_sheet(workbook["三大件专项"])
    standard_sections.append(power)
    summaries.append(summary)
    new_energy, summary = build_new_energy_sheet(workbook["新能源专项"])
    summaries.append(summary)
    rules, excluded = build_rules(workbook["判定总则"])
    total_rows = sum(item["validRows"] for item in summaries)
    total_items = sum(item["checkItems"] for item in summaries)
    total_criteria = sum(item["criteria"] for item in summaries)
    source_manifest = {
        "workbook": workbook_path.name,
        "sheets": ["整车通用", "左前方位", "左后方位", "右前方位", "右后方位", "三大件专项", "新能源专项", "判定总则"],
        "sheetSummaries": summaries,
        "rules": rules,
        "excludedRows": excluded,
        "totals": {"applicableSourceRows": total_rows, "checkItems": total_items, "criteria": total_criteria, "activeRules": len(rules), "excludedRows": len(excluded)},
    }
    return {
        "sourceManifest": source_manifest,
        "standard": {"code": "STANDARD_VEHICLE", "name": "乘用车通用鉴定模板", "description": "完整导入整车通用、四方位、三大件专项；新能源车辆使用独立同结构版本。", "energyTypeScope": "ICE", "sections": standard_sections},
        "newEnergy": {"code": "NEW_ENERGY_VEHICLE", "name": "新能源通用鉴定模板", "description": "完整复用通用鉴定维度并追加新能源专项；所有新能源车辆统一使用本模板。", "energyTypeScope": "NEW_ENERGY", "sections": standard_sections + [new_energy]},
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    catalog = build_catalog(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    manifest = catalog["sourceManifest"]
    print(json.dumps({"totals": manifest["totals"], "sheets": [{"sheet": item["sheet"], "validRows": item["validRows"], "checkItems": item["checkItems"], "criteria": item["criteria"], "fallbackCriteria": item["fallbackCriteria"]} for item in manifest["sheetSummaries"]], "rules": manifest["rules"], "excludedRows": manifest["excludedRows"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
